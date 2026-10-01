"""
Organizador de vagas a partir dos e-mails de alerta (LinkedIn, Jobbol etc.).

O que faz:
- Lê o Gmail via IMAP (somente leitura; não altera nada no Gmail nem no LinkedIn)
- Extrai TODAS as vagas de cada e-mail (cargo, empresa, local, salário, link)
- Ignora e-mails que não são vagas (confirmação de candidatura, pesquisas etc.)
- Remove repetidas pelo número da vaga e dá uma nota de compatibilidade
- Gera/atualiza vagas.xlsx com filtros, links clicáveis e coluna de Status.
  Ao rodar de novo, o que você marcou (Status e Notas) é mantido.

Como usar:
1. Crie um arquivo .env ao lado deste script:
       GMAIL_USER=seuemail@gmail.com
       GMAIL_APP_PASSWORD=xxxx xxxx xxxx xxxx
2. pip install python-dotenv openpyxl
3. Feche o vagas.xlsx se estiver aberto no Excel e rode:
       python organizador_vagas.py
"""

import email
import hashlib
import html
import imaplib
import os
import re
from datetime import date, datetime, timedelta
from email.header import decode_header, make_header
from email.utils import parsedate_to_datetime
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

DIAS = 30
LIMITE_NOTA = -3  # vagas com nota igual ou menor que isso nem entram na planilha
ARQUIVO = Path("vagas.xlsx")

REMETENTES = ["jobalerts-noreply@linkedin.com", "jobs-noreply@linkedin.com", "jobbol.com.br"]

# E-mails com estes trechos no assunto não são vagas.
IGNORAR_ASSUNTO = [
    "candidatura", "alerta de vaga", "pesquisa", "premium", "compartilham novidades",
    "impressões", "competência", "está em alta", "pós-graduação", "saiba mais sobre",
]

# A vaga só entra se o cargo parecer da área de desenvolvimento.
REGEX_DEV = re.compile(
    r"desenvolv|developer|\bdev\b|programador|engenh\w* de software|software engineer|"
    r"front|back|full ?stack|analista de sistemas|analista de desenvolvimento|\bweb\b",
    re.I,
)

# (expressão, pontos). Ajuste ao seu gosto.
PESOS = [
    (r"front[- ]?end", 3), (r"react|angular|\bvue\b", 2), (r"javascript|typescript|\bjs\b|\bts\b", 2),
    (r"\bhtml\b|\bcss\b", 1), (r"j[uú]nior|\bjr\b", 2), (r"remot", 2), (r"python", 1),
    (r"\bia\b|intelig[eê]ncia artificial|\bllm\b", 1), (r"full ?stack", 1),
    (r"s[eê]nior|\bsr\b", -6), (r"pleno|\bpl\b", -5),
    (r"especialista|gerente|coordenador|\blead\b", -6),
    (r"\bjava\b", -5), (r"\.net|c#|\bphp\b|\bruby\b|\bgolang\b", -3),
]

STATUS = ["Nova", "Salva", "Candidatado", "Entrevista", "Oferta", "Descartada"]
COLUNAS = ["Nota", "Status", "Cargo", "Empresa", "Local", "Remoto", "Salário", "Data", "Link", "Notas", "ID"]

RE_ID = re.compile(r"jobs/view/(\d+)")
RE_SALARIO = re.compile(r"R\$\s?[\d\.,]+\s?(?:mil)?(?:\s?[-–]\s?R\$\s?[\d\.,]+\s?(?:mil)?)?")


# ------------------------------------------------------------------ leitura

def decodificar(texto):
    return str(make_header(decode_header(texto or "")))


def texto_do_email(msg):
    """Prefere a parte text/plain; se não houver, converte o HTML em texto."""
    html_bruto = ""
    for parte in msg.walk():
        tipo = parte.get_content_type()
        if tipo not in ("text/plain", "text/html"):
            continue
        charset = parte.get_content_charset() or "utf-8"
        conteudo = parte.get_payload(decode=True).decode(charset, errors="ignore")
        if tipo == "text/plain":
            return conteudo
        html_bruto = conteudo
    texto = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>|</td>", "\n", html_bruto)
    texto = re.sub(r"(?i)<a [^>]*href=\"([^\"]+)\"[^>]*>", r" \1 ", texto)
    return html.unescape(re.sub(r"<[^>]+>", "", texto))


def vagas_do_corpo(texto):
    """Cada vaga aparece como: Cargo / Empresa / Local / 'Visualizar vaga: <link>'."""
    linhas = [l.strip() for l in texto.splitlines()]
    vagas = []
    for i, linha in enumerate(linhas):
        achou = RE_ID.search(linha)
        if not achou or "http" not in linha:
            continue
        previas, j = [], i - 1
        while j >= 0 and len(previas) < 3:
            atual = linhas[j]
            if atual.startswith("---") or "http" in atual:
                break
            if atual:
                previas.append(atual)
            j -= 1
        if len(previas) == 3:
            local, empresa, cargo = previas
            vagas.append({
                "id": achou.group(1), "cargo": cargo, "empresa": empresa, "local": local,
                "link": f"https://www.linkedin.com/jobs/view/{achou.group(1)}/",
            })
    return vagas


def vaga_do_assunto(assunto):
    """Plano B: assuntos do tipo '<Cargo> na empresa <Empresa>'."""
    if " na empresa " not in assunto:
        return []
    cargo, empresa = assunto.rsplit(" na empresa ", 1)
    chave = hashlib.md5(assunto.encode()).hexdigest()[:10]
    return [{"id": f"s{chave}", "cargo": cargo.strip(), "empresa": empresa.strip(), "local": "", "link": ""}]


def eh_remoto(vaga):
    texto = f"{vaga['cargo']} {vaga['local']}".lower()
    return "Sim" if re.search(r"remot|home ?office|anywhere", texto) else ""


def pontuar(vaga):
    texto = f"{vaga['cargo']} {vaga['empresa']} {vaga['local']}".lower()
    return sum(p for padrao, p in PESOS if re.search(padrao, texto))


def buscar_vagas(usuario, senha):
    desde = (date.today() - timedelta(days=DIAS)).strftime("%d-%b-%Y")
    vagas = {}
    with imaplib.IMAP4_SSL("imap.gmail.com") as imap:
        imap.login(usuario, senha)
        imap.select("INBOX", readonly=True)  # readonly: não marca nada como lido
        ids = set()
        for remetente in REMETENTES:
            _, dados = imap.search(None, f'(SINCE "{desde}" FROM "{remetente}")')
            ids.update(dados[0].split())

        print(f"{len(ids)} e-mails de vagas encontrados. Lendo...")
        for n, i in enumerate(sorted(ids), start=1):
            if n % 10 == 0:
                print(f"  {n}/{len(ids)}")
            _, dados = imap.fetch(i, "(RFC822)")
            msg = email.message_from_bytes(dados[0][1])
            assunto = decodificar(msg["Subject"])
            if any(t in assunto.lower() for t in IGNORAR_ASSUNTO):
                continue
            data_email = parsedate_to_datetime(msg["Date"]).strftime("%Y-%m-%d")
            encontradas = vagas_do_corpo(texto_do_email(msg)) or vaga_do_assunto(assunto)
            for v in encontradas:
                if v["id"] in vagas or not REGEX_DEV.search(v["cargo"]):
                    continue
                salario = RE_SALARIO.search(f"{assunto} {v['cargo']}")
                v.update(data=data_email, salario=salario.group(0) if salario else "")
                vagas[v["id"]] = v
    return vagas


# ------------------------------------------------------------------ planilha

def carregar_existentes():
    """Lê a planilha atual para não perder Status e Notas que você já marcou."""
    if not ARQUIVO.exists():
        return {}
    ws = load_workbook(ARQUIVO)["Vagas"]
    existentes = {}
    cabecalho = [c.value for c in ws[1]]
    for linha in ws.iter_rows(min_row=2):
        d = dict(zip(cabecalho, [c.value for c in linha]))
        if not d.get("ID"):
            continue
        celula_link = linha[cabecalho.index("Link")]
        link = celula_link.hyperlink.target if celula_link.hyperlink else ""
        existentes[str(d["ID"])] = {
            "id": str(d["ID"]), "cargo": d["Cargo"] or "", "empresa": d["Empresa"] or "",
            "local": d["Local"] or "", "salario": d["Salário"] or "", "data": d["Data"] or "",
            "link": link, "status": d["Status"] or "Nova", "notas": d["Notas"] or "",
        }
    return existentes


def escrever_planilha(vagas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Vagas"
    fonte = "Arial"

    ws.append(COLUNAS)
    for c in ws[1]:
        c.font = Font(name=fonte, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    ordenadas = sorted(vagas.values(), key=lambda v: (v["nota"], v["data"]), reverse=True)
    for n, v in enumerate(ordenadas, start=2):
        ws.append([v["nota"], v.get("status", "Nova"), v["cargo"], v["empresa"], v["local"],
                   v.get("remoto", ""), v["salario"], v["data"], "Abrir vaga" if v["link"] else "",
                   v.get("notas", ""), v["id"]])
        for c in ws[n]:
            c.font = Font(name=fonte, size=10)
            c.alignment = Alignment(vertical="center", wrap_text=c.column_letter in "CDEJ")
        for col in (1, 2, 6):  # Nota, Status e Remoto centralizados
            ws.cell(n, col).alignment = Alignment(horizontal="center", vertical="center")
        if v["link"]:
            col_link = COLUNAS.index("Link") + 1
            ws.cell(n, col_link).hyperlink = v["link"]
            ws.cell(n, col_link).font = Font(name=fonte, size=10, color="0563C1", underline="single")

    ultima = max(ws.max_row, 2)
    for letra, largura in zip("ABCDEFGHIJK", [9, 14, 46, 26, 22, 10, 16, 12, 12, 32, 12]):
        ws.column_dimensions[letra].width = largura
    ws.column_dimensions["K"].hidden = True
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:K{ultima}"

    # lista suspensa de status
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B2:B{max(ultima, 500)}")

    # cores: nota (verde = melhor) e status
    ws.conditional_formatting.add(
        f"A2:A{max(ultima, 500)}",
        ColorScaleRule(start_type="num", start_value=0, start_color="F8F8F8",
                       end_type="num", end_value=8, end_color="63BE7B"))
    cores = {"Candidatado": "BDD7EE", "Entrevista": "FFE699", "Oferta": "A9D08E", "Descartada": "D9D9D9"}
    for status, cor in cores.items():
        ws.conditional_formatting.add(
            f"B2:B{max(ultima, 500)}",
            CellIsRule(operator="equal", formula=[f'"{status}"'], fill=PatternFill("solid", bgColor=cor)))

    # aba de resumo com fórmulas (atualiza sozinha quando você muda o Status)
    rs = wb.create_sheet("Resumo")
    rs["A1"], rs["A1"].font = "Resumo das vagas", Font(name=fonte, bold=True, size=14)
    rs.append([])
    linhas = [("Total de vagas", "=COUNTA(Vagas!C2:C5000)")]
    linhas += [(s, f'=COUNTIF(Vagas!B2:B5000,"{s}")') for s in STATUS]
    linhas += [("Nota 5 ou mais", '=COUNTIF(Vagas!A2:A5000,">=5")')]
    for nome, formula in linhas:
        rs.append([nome, formula])
    rs.append([])
    rs.append(["Atualizado em", datetime.now().strftime("%d/%m/%Y %H:%M")])
    rs.append([])
    rs.append(["Como usar"])
    rs.append(["1. Mude o Status na aba Vagas (lista suspensa) conforme avança em cada vaga."])
    rs.append(["2. Use os filtros do cabeçalho para ver só o que interessa (ex.: Status = Nova)."])
    rs.append(["3. Escreva observações na coluna Notas. Ao rodar o script de novo, nada disso se perde."])
    for linha in rs.iter_rows():
        for c in linha:
            if c.font != Font(name=fonte, bold=True, size=14):
                c.font = Font(name=fonte, size=10, bold=(c.column == 1 and c.row in (2, 10, 13)))
    rs["A1"].font = Font(name=fonte, bold=True, size=14)
    rs.column_dimensions["A"].width = 30
    rs.column_dimensions["B"].width = 18

    wb.save(ARQUIVO)


def main():
    load_dotenv()
    usuario, senha = os.environ["GMAIL_USER"], os.environ["GMAIL_APP_PASSWORD"]

    existentes = carregar_existentes()
    novas = buscar_vagas(usuario, senha)

    total_novas = 0
    for id_, v in novas.items():
        if id_ in existentes:
            existentes[id_].update(data=max(existentes[id_]["data"], v["data"]))
        elif pontuar(v) > LIMITE_NOTA:  # descarta vagas claramente acima do seu nível
            v.update(status="Nova", notas="")
            existentes[id_] = v
            total_novas += 1
    for v in existentes.values():
        v["nota"] = pontuar(v)
        v["remoto"] = eh_remoto(v)
    # limpa vagas antigas que ficaram abaixo do seu nível (só se você não mexeu no Status)
    existentes = {k: v for k, v in existentes.items()
                  if v["nota"] > LIMITE_NOTA or v.get("status", "Nova") != "Nova"}

    try:
        escrever_planilha(existentes)
    except PermissionError:
        print(f"Feche o {ARQUIVO} no Excel e rode de novo.")
        return

    print(f"{total_novas} vagas novas. Total na planilha: {len(existentes)} ({ARQUIVO})")
    for v in sorted(existentes.values(), key=lambda v: v["nota"], reverse=True)[:5]:
        print(f"  [{v['nota']:>2}] {v['cargo']} — {v['empresa']}")


if __name__ == "__main__":
    main()
